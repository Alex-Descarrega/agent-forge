"""Paquete de agente: lo que se afina y lo que se traslada de entorno.

Layout de un workspace de entrenamiento:

    <workspace>/
      spec.json              # salida de la entrevista
      samples/               # muestras reales que aportas (txt/md/json)
      evals/cases.json       # batería de tests (con split fit/validation)
      package/               # versión actual del agente  <-- esto es lo portátil
        agent.md             # frontmatter + system prompt (formato subagente Claude Code)
        tools.json           # contrato de herramientas (abstracto)
        config.json          # modelo, max_turns, etc.
        skills/<name>/SKILL.md
      versions/v001/...      # snapshots de package/ en cada iteración aceptada
      runs/v001-fit.json     # informes de evaluación
      report.md              # informe final
"""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .schemas import AgentSpec, ToolContract, dump_json, load_json


AGENT_MD_TEMPLATE = """---
name: {name}
description: {description}
tools: {tools}
model: {model}
---

{system_prompt}
"""


@dataclass
class AgentPackage:
    name: str
    description: str
    system_prompt: str
    tools: list[ToolContract] = field(default_factory=list)
    skills: dict[str, str] = field(default_factory=dict)      # nombre -> contenido SKILL.md
    config: dict[str, Any] = field(default_factory=lambda: {"model": "sonnet", "max_turns": 12})

    # ---------------- carga / guardado ----------------

    @classmethod
    def load(cls, package_dir: Path) -> "AgentPackage":
        agent_md = (package_dir / "agent.md").read_text(encoding="utf-8")
        fm, body = _split_frontmatter(agent_md)
        tools_path = package_dir / "tools.json"
        tools = [ToolContract.from_dict(t) for t in load_json(tools_path)] if tools_path.exists() else []
        config_path = package_dir / "config.json"
        config = load_json(config_path) if config_path.exists() else {"model": fm.get("model", "sonnet"), "max_turns": 12}
        skills: dict[str, str] = {}
        skills_dir = package_dir / "skills"
        if skills_dir.exists():
            for sk in sorted(skills_dir.iterdir()):
                f = sk / "SKILL.md"
                if f.exists():
                    skills[sk.name] = f.read_text(encoding="utf-8")
        return cls(
            name=fm.get("name", package_dir.parent.name),
            description=fm.get("description", ""),
            system_prompt=body.strip(),
            tools=tools,
            skills=skills,
            config=config,
        )

    def save(self, package_dir: Path) -> None:
        package_dir.mkdir(parents=True, exist_ok=True)
        (package_dir / "agent.md").write_text(self.render_agent_md(), encoding="utf-8")
        dump_json(package_dir / "tools.json", [t.__dict__ for t in self.tools])
        dump_json(package_dir / "config.json", self.config)
        skills_dir = package_dir / "skills"
        if skills_dir.exists():
            shutil.rmtree(skills_dir)
        for name, content in self.skills.items():
            d = skills_dir / name
            d.mkdir(parents=True, exist_ok=True)
            (d / "SKILL.md").write_text(content, encoding="utf-8")

    def render_agent_md(self) -> str:
        """Frontmatter válido para un subagente de Claude Code.

        Claude Code espera los nombres de tools MCP como `mcp__<servidor>__<tool>`; el nombre del
        servidor es del entorno destino, así que se toma de config["mcp_server"] (por defecto "forge").
        """
        server = self.config.get("mcp_server", "forge")
        tools = ", ".join(f"mcp__{server}__{t.name}" for t in self.tools) or "Read"
        return AGENT_MD_TEMPLATE.format(
            name=self.name,
            description=json.dumps(self.description.replace("\n", " "), ensure_ascii=False),  # entrecomillado YAML-safe
            tools=tools,
            model=self.config.get("model", "sonnet"),
            system_prompt=self.system_prompt.strip(),
        )

    # ---------------- fábrica desde la spec ----------------

    @classmethod
    def initial_from_spec(cls, spec: AgentSpec) -> "AgentPackage":
        """Versión 0: un system prompt honesto pero genérico. El bucle lo irá afinando."""
        lines = [f"Eres **{spec.name}**, un agente de tipo {spec.kind}.", "", f"## Objetivo", spec.objective, ""]
        if spec.audience:
            lines += ["## Con quién trabajas", spec.audience, ""]
        if spec.success_criteria:
            lines += ["## Qué significa hacerlo bien"] + [f"- {c}" for c in spec.success_criteria] + [""]
        if spec.never_do:
            lines += ["## Nunca"] + [f"- {c}" for c in spec.never_do] + [""]
        if spec.tone:
            lines += ["## Tono", spec.tone, ""]
        if spec.tools:
            lines += ["## Herramientas"] + [f"- `{t.name}`: {t.description}" for t in spec.tools] + [""]
        lines += ["## Método", "1. Lee el contexto completo antes de actuar.",
                  "2. Usa las herramientas cuando necesites información o cuando debas actuar; no inventes datos.",
                  "3. Si falta información imprescindible, pídela antes de actuar.",
                  "4. Explica brevemente qué has hecho al terminar."]
        skills = {}
        if spec.methodology.strip():
            skills["metodologia"] = (
                "---\nname: metodologia\ndescription: Cómo trabajamos en este dominio. Cárgala antes de resolver cualquier caso.\n---\n\n"
                + spec.methodology.strip() + "\n"
            )
        return cls(
            name=spec.name,
            description=f"{spec.kind}: {spec.objective}",
            system_prompt="\n".join(lines),
            tools=list(spec.tools),
            skills=skills,
        )

    # ---------------- aplicar parches del editor ----------------

    def apply_patch(self, patch: dict[str, Any]) -> "AgentPackage":
        """Devuelve una copia con los cambios propuestos por el editor.

        Formato de patch (todo opcional):
          {"system_prompt": "...texto completo nuevo...",
           "skills": {"nombre": "...contenido completo..." | null para borrar},
           "config": {"max_turns": 8},
           "rationale": "por qué"}
        """
        new = AgentPackage(
            name=self.name,
            description=patch.get("description", self.description),
            system_prompt=patch.get("system_prompt", self.system_prompt),
            tools=list(self.tools),
            skills=dict(self.skills),
            config={**self.config, **(patch.get("config") or {})},
        )
        for k, v in (patch.get("skills") or {}).items():
            if v is None:
                new.skills.pop(k, None)
            else:
                new.skills[k] = v
        return new

    def __eq__(self, other: object) -> bool:
        if not isinstance(other, AgentPackage):
            return NotImplemented
        return (self.system_prompt.strip() == other.system_prompt.strip() and self.skills == other.skills
                and self.config == other.config and self.description == other.description)

    def summary(self) -> str:
        return (
            f"{self.name} | prompt {len(self.system_prompt)} chars | "
            f"{len(self.skills)} skills | {len(self.tools)} tools | model {self.config.get('model')}"
        )


# ---------------------------------------------------------------------------
# Workspace
# ---------------------------------------------------------------------------

class Workspace:
    def __init__(self, root: Path):
        self.root = Path(root)

    @property
    def spec_path(self) -> Path: return self.root / "spec.json"
    @property
    def samples_dir(self) -> Path: return self.root / "samples"
    @property
    def cases_path(self) -> Path: return self.root / "evals" / "cases.json"
    @property
    def package_dir(self) -> Path: return self.root / "package"
    @property
    def versions_dir(self) -> Path: return self.root / "versions"
    @property
    def runs_dir(self) -> Path: return self.root / "runs"

    def init(self) -> None:
        for d in (self.samples_dir, self.cases_path.parent, self.package_dir, self.versions_dir, self.runs_dir):
            d.mkdir(parents=True, exist_ok=True)
        readme = self.samples_dir / "README.md"
        if not readme.exists():
            readme.write_text(
                "Deja aquí muestras reales: tickets, conversaciones, respuestas buenas y malas.\n"
                "Formatos: .txt, .md, .json, .html, .pptx; también .docx/.odt/.rtf (con pandoc) y .pdf (con pdftotext o pypdf). Un archivo = una muestra.\n"
                "Si marcas un archivo como *.good.md o *.bad.md, el constructor de evals lo usará como ejemplo positivo/negativo.\n",
                encoding="utf-8",
            )

    def load_spec(self) -> AgentSpec:
        return AgentSpec.from_dict(load_json(self.spec_path))

    def save_spec(self, spec: AgentSpec) -> None:
        dump_json(self.spec_path, spec.to_dict())

    SAMPLE_TEXT_EXT = {".txt", ".md", ".json", ".csv", ".html", ".htm"}
    SAMPLE_DOC_EXT = {".docx", ".pdf", ".odt", ".rtf", ".pptx"}

    def load_samples(self) -> list[dict[str, str]]:
        """Lee samples/. Los .docx/.pdf/.odt/.rtf/.pptx se convierten a texto una vez y se cachean
        en samples/.text/<nombre>.txt (hace falta `pandoc` para docx/odt/rtf, `pdftotext` para pdf,
        o python-docx / pypdf si están instalados). Lo que no se pueda convertir se lista con aviso."""
        out = []
        if not self.samples_dir.exists():
            return out
        skipped = []
        for f in sorted(self.samples_dir.iterdir()):
            if f.name == "README.md" or f.name.startswith(".") or f.is_dir():
                continue
            label = "good" if ".good." in f.name else "bad" if ".bad." in f.name else "neutral"
            if f.suffix.lower() in self.SAMPLE_TEXT_EXT:
                try:
                    out.append({"file": f.name, "label": label, "content": f.read_text(encoding="utf-8", errors="replace")})
                except OSError:
                    skipped.append(f.name)
            elif f.suffix.lower() in self.SAMPLE_DOC_EXT:
                text = _doc_to_text(f, self.samples_dir / ".text")
                if text:
                    out.append({"file": f.name, "label": label, "content": text})
                else:
                    skipped.append(f.name)
            else:
                skipped.append(f.name)
        # dedup: copias "(2)", "(3)" o contenido idéntico
        import hashlib
        seen_hash: dict[str, str] = {}
        deduped = []
        for smp in out:
            h = hashlib.sha1(smp["content"].encode("utf-8", errors="replace")).hexdigest()
            if h in seen_hash:
                continue
            seen_hash[h] = smp["file"]
            deduped.append(smp)
        if len(deduped) < len(out):
            print(f"[samples] {len(out) - len(deduped)} duplicados ignorados (mismo contenido); {len(deduped)} muestras únicas")
        out = deduped
        if skipped:
            print(f"[samples] ignorados (formato no soportado o sin conversor instalado): {', '.join(skipped[:8])}"
                  + (" …" if len(skipped) > 8 else "") + "\n          Para .docx/.odt: instala pandoc; para .pdf: poppler-utils (pdftotext) o pip install pypdf; .pptx se lee sin dependencias (mejor con pip install python-pptx).")
        return out

    def load_package(self) -> AgentPackage:
        return AgentPackage.load(self.package_dir)

    def save_package(self, pkg: AgentPackage) -> None:
        pkg.save(self.package_dir)

    def next_version(self) -> str:
        existing = sorted(p.name for p in self.versions_dir.glob("v*"))
        n = int(existing[-1][1:]) + 1 if existing else 1
        return f"v{n:03d}"

    def snapshot(self, pkg: AgentPackage, version: str, meta: dict[str, Any] | None = None) -> Path:
        d = self.versions_dir / version
        pkg.save(d)
        if meta:
            dump_json(d / "meta.json", meta)
        return d

    def export(self, out_zip: Path) -> Path:
        """Empaqueta package/ + evals/ + spec.json para trasladarlo a otro entorno."""
        staging = self.root / "_export"
        if staging.exists():
            shutil.rmtree(staging)
        staging.mkdir()
        shutil.copytree(self.package_dir, staging / "package")
        if self.cases_path.exists():
            shutil.copytree(self.cases_path.parent, staging / "evals")
        if self.spec_path.exists():
            shutil.copy(self.spec_path, staging / "spec.json")
        if (self.root / "report.md").exists():
            shutil.copy(self.root / "report.md", staging / "report.md")
        (staging / "README.md").write_text(_EXPORT_README, encoding="utf-8")
        out = shutil.make_archive(str(out_zip.with_suffix("")), "zip", staging)
        shutil.rmtree(staging)
        return Path(out)


_EXPORT_README = """# Paquete de agente exportado

- `package/agent.md`   -> cópialo a `.claude/agents/<nombre>.md` (Claude Code) o usa su cuerpo como `system_prompt` (Agent SDK / API).
                          El campo `tools:` usa el prefijo `mcp__<servidor>__`; pon el nombre real de tu servidor MCP en
                          `config.json` -> "mcp_server" antes de exportar, o edítalo a mano.
- `package/skills/*`   -> cópialos a `.claude/skills/` o al directorio de skills de tu plugin.
- `package/tools.json` -> contrato de herramientas. En el entorno destino, cada tool debe existir
                          (como conector MCP o tool propia) con el mismo nombre y entrada.
- `evals/cases.json`   -> vuelve a ejecutar `forge run` en el entorno destino para comprobar que nada se ha roto.
- `spec.json`          -> la especificación original, por si hay que reentrevistar.
"""


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _pptx_to_text(f: Path) -> str | None:
    """Texto de una presentación, diapositiva a diapositiva, con tablas y notas.
    Usa python-pptx si está; si no, lee el XML del zip directamente (sin dependencias)."""
    import re as _re
    import zipfile
    try:
        from pptx import Presentation  # type: ignore
        from pptx.util import Emu  # noqa: F401
        prs = Presentation(str(f))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            out.append(f"\n## Diapositiva {i}\n")
            shapes = sorted(slide.shapes, key=lambda sh: ((sh.top or 0), (sh.left or 0)))
            for sh in shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    out.append(sh.text_frame.text.strip())
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        out.append(" | ".join(c.text.strip() for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                out.append("[notas] " + slide.notes_slide.notes_text_frame.text.strip())
        return "\n".join(out)
    except ImportError:
        pass
    except Exception:  # noqa: BLE001
        return None
    # fallback sin dependencias: XML crudo de cada slide en orden
    try:
        with zipfile.ZipFile(f) as z:
            names = sorted((n for n in z.namelist() if _re.match(r"ppt/slides/slide\d+\.xml$", n)),
                           key=lambda n: int(_re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
            out = []
            for i, n in enumerate(names, 1):
                xml = z.read(n).decode("utf-8", errors="replace")
                paras = []
                for p in _re.findall(r"<a:p>(.*?)</a:p>", xml, _re.S):
                    runs = _re.findall(r"<a:t>(.*?)</a:t>", p, _re.S)
                    if runs:
                        paras.append("".join(runs))
                out.append(f"\n## Diapositiva {i}\n" + "\n".join(paras))
            text = "\n".join(out)
            text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'")
            return text if text.strip() else None
    except Exception:  # noqa: BLE001
        return None


def _doc_to_text(f: Path, cache_dir: Path) -> str | None:
    """Convierte un documento a texto plano con lo que haya instalado; cachea por mtime."""
    import shutil as _sh
    import subprocess
    cache_dir.mkdir(parents=True, exist_ok=True)
    cached = cache_dir / (f.name + ".txt")
    if cached.exists() and cached.stat().st_mtime >= f.stat().st_mtime:
        return cached.read_text(encoding="utf-8")
    text: str | None = None
    ext = f.suffix.lower()
    try:
        if ext == ".pdf":
            if _sh.which("pdftotext"):
                r = subprocess.run(["pdftotext", "-layout", str(f), "-"], capture_output=True, text=True, timeout=120)
                text = r.stdout if r.returncode == 0 else None
            if not text:
                try:
                    from pypdf import PdfReader  # type: ignore
                    text = "\n".join((p.extract_text() or "") for p in PdfReader(str(f)).pages)
                except ImportError:
                    pass
        elif ext == ".pptx":
            text = _pptx_to_text(f)
        elif ext in {".docx", ".odt", ".rtf"}:
            if _sh.which("pandoc"):
                r = subprocess.run(["pandoc", str(f), "-t", "gfm", "--wrap=none"], capture_output=True, text=True, timeout=120)
                text = r.stdout if r.returncode == 0 else None
            if not text and ext == ".docx":
                try:
                    import docx  # type: ignore
                    d = docx.Document(str(f))
                    parts = [p.text for p in d.paragraphs]
                    for t in d.tables:
                        for row in t.rows:
                            parts.append(" | ".join(c.text.strip() for c in row.cells))
                    text = "\n".join(parts)
                except ImportError:
                    pass
    except Exception:  # noqa: BLE001
        text = None
    if text and text.strip():
        cached.write_text(text, encoding="utf-8")
        return text
    return None


_FM_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n(.*)$", re.S)


def _split_frontmatter(text: str) -> tuple[dict[str, str], str]:
    m = _FM_RE.match(text)
    if not m:
        return {}, text
    fm: dict[str, str] = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            v = v.strip()
            if len(v) >= 2 and v[0] == v[-1] and v[0] in "\"'":
                try:
                    v = json.loads(v) if v[0] == '"' else v[1:-1]
                except json.JSONDecodeError:
                    v = v[1:-1]
            fm[k.strip()] = v
    return fm, m.group(2)
